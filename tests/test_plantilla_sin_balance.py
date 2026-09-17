"""El balance de prueba es opcional, y el papel tiene que poder escribirse.

Salio de produccion: una revision sin balance reventaba al DEPOSITAR, no al
revisar. El motor ya era disciplinado -- ctx.saldos queda en None y C2 sale
NO_EJECUTADO -- pero la hoja REVISION ICA hacia ctx.saldos.get(...) a secas y
el auditor recibia un 500 despues de que el motor habia hecho todo el trabajo.

La regla que fija este archivo es la misma que ya estaba escrita en esa
funcion para el caso de la cuenta faltante: si el dato no esta, el papel lo
DICE. Un cero mudo en esa celda seria justo lo que este proyecto no acepta.
"""

import shutil
import warnings
from pathlib import Path

import openpyxl
import pytest

from motor_reteica.parametros.municipios.santa_marta import MUNICIPIO
from motor_reteica.pipeline import revisar
from motor_reteica.plantilla.deposito import depositar
from motor_reteica.tipos import Estado

BASE = Path(__file__).parent / "fixtures" / "terlica_202607"


@pytest.fixture(scope="module")
def sin_balance(tmp_path_factory):
    """Las fixtures completas, menos el balance de prueba."""
    carpeta = tmp_path_factory.mktemp("sin_balance") / "corrida"
    shutil.copytree(BASE, carpeta)
    (carpeta / "balance.xlsx").unlink()
    return revisar(carpeta, nit="819002433", periodo="2026-07",
                   municipio=MUNICIPIO)


@pytest.fixture(scope="module")
def papel(sin_balance, tmp_path_factory):
    salida = tmp_path_factory.mktemp("papel") / "papel.xlsx"
    depositar(sin_balance, salida)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(salida)


def test_sin_balance_el_motor_deja_saldos_en_none(sin_balance):
    """La premisa: no es que el balance salga vacio, es que no hay balance."""
    assert sin_balance.saldos is None


def test_c2_queda_no_ejecutado_no_ok(sin_balance):
    c2 = next(r for r in sin_balance.resultados if r.codigo == "C2")
    assert c2.estado is Estado.NO_EJECUTADO


def test_el_papel_se_escribe_sin_balance(papel):
    """Lo que fallaba: depositar() lanzaba AttributeError y salia un 500."""
    assert "REVISION ICA" in papel.sheetnames


def test_la_celda_dice_que_falta_el_balance_en_vez_de_callarlo(papel):
    hoja = papel["REVISION ICA"]
    for celda in ("N12", "N13"):
        valor = hoja[celda].value
        assert isinstance(valor, str), (
            "%s trae %r: un numero ahi afirma un saldo que nadie aporto"
            % (celda, valor))
        assert "BALANCE" in valor.upper()


def test_con_balance_la_celda_sigue_trayendo_el_saldo(tmp_path_factory):
    """La guarda nueva no puede haber roto el camino normal."""
    ctx = revisar(BASE, nit="819002433", periodo="2026-07",
                  municipio=MUNICIPIO)
    salida = tmp_path_factory.mktemp("con") / "papel.xlsx"
    depositar(ctx, salida)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        libro = openpyxl.load_workbook(salida)
    assert isinstance(libro["REVISION ICA"]["N13"].value, int)
