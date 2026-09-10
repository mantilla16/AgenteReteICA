"""La procedencia declarada de una fuente tiene que LLEGAR AL PAPEL.

Bug real encontrado en revision: el manifiesto de la carpeta de julio trae un
campo _nota_balance explicando que el balance NO lo entrego el cliente, sino
que se extrajo de la hoja BALANCE del papel de trabajo MANUAL DE LA FIRMA.
La justificacion es defendible -- el contenido es un export crudo de SAP y las
cifras coinciden -- pero el parser descartaba el campo y el socio leia en la
hoja Parametros:

    balance | balance_terlica_202607.xlsx | <sha256>

sin una palabra sobre de donde salio. El papel mostraba un OK de "el auxiliar
cuadra contra el balance" donde el balance lo produjo la propia firma.

Es exactamente la leccion que M15 escribio para si mismo: verificar_tipo_
documento comprueba ESTRUCTURA, no PROCEDENCIA. La procedencia solo la puede
declarar quien arma el manifiesto, y si no llega al papel no sirve de nada.
"""

import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from motor_reteica.manifiesto import leer_manifiesto
from motor_reteica.papel_excel import generar_papel
from motor_reteica.parametros.municipios.santa_marta import MUNICIPIO
from motor_reteica.pipeline import revisar

BASE = Path(__file__).parent / "fixtures" / "terlica_202607"

_NOTA = ("El cliente NO entrega balance como archivo independiente: se extrajo "
         "de la hoja BALANCE del papel de trabajo manual de la firma.")


def _carpeta_con_manifiesto(tmp_path, notas=None):
    destino = tmp_path / "corrida"
    shutil.copytree(BASE, destino)
    contenido = {
        "nit": "819002433", "periodo": "2026-07", "municipio": "Santa Marta",
        "declarado_por": "analitica@rbcol.co", "fecha": "2026-09-10",
        "archivos": {
            "borrador": "borrador.pdf", "auxiliar": "auxiliar_2368.xlsx",
            "balance": "balance.xlsx", "erp": "sap_retenciones.xlsx",
            "pago_anterior": None, "facturas": "facturas",
        },
    }
    contenido.update(notas or {})
    (destino / "manifiesto.json").write_text(
        json.dumps(contenido, ensure_ascii=False), encoding="utf-8")
    return destino


def test_el_parser_conserva_las_notas_de_procedencia(tmp_path):
    """Antes se descartaban en silencio al construir el Manifiesto."""
    carpeta = _carpeta_con_manifiesto(tmp_path, {"_nota_balance": _NOTA})
    manifiesto = leer_manifiesto(carpeta)
    assert manifiesto.notas.get("balance") == _NOTA


def test_una_nota_de_cualquier_rol_se_conserva(tmp_path):
    """El mecanismo es general: no esta cableado al balance."""
    carpeta = _carpeta_con_manifiesto(
        tmp_path, {"_nota_erp": "export regenerado el 11.08.2026"})
    assert leer_manifiesto(carpeta).notas.get("erp") == \
        "export regenerado el 11.08.2026"


def test_sin_notas_el_manifiesto_sigue_siendo_valido(tmp_path):
    assert leer_manifiesto(_carpeta_con_manifiesto(tmp_path)).notas == {}


def test_la_procedencia_se_imprime_en_el_papel(tmp_path):
    """Lo que de verdad importa: que el socio la lea."""
    carpeta = _carpeta_con_manifiesto(tmp_path, {"_nota_balance": _NOTA})
    ctx = revisar(carpeta, nit="819002433", periodo="2026-07",
                  municipio=MUNICIPIO)
    salida = generar_papel(carpeta / "papel.xlsx", ctx)

    hoja = openpyxl.load_workbook(salida)["Parametros"]
    texto = " ".join(str(c.value) for fila in hoja.iter_rows()
                     for c in fila if c.value is not None)
    assert "PROCEDENCIA" in texto.upper()
    assert "papel de trabajo manual de la firma" in texto


def test_sin_nota_no_aparece_un_bloque_vacio(tmp_path):
    carpeta = _carpeta_con_manifiesto(tmp_path)
    ctx = revisar(carpeta, nit="819002433", periodo="2026-07",
                  municipio=MUNICIPIO)
    salida = generar_papel(carpeta / "papel.xlsx", ctx)
    hoja = openpyxl.load_workbook(salida)["Parametros"]
    texto = " ".join(str(c.value) for fila in hoja.iter_rows()
                     for c in fila if c.value is not None)
    assert "PROCEDENCIA" not in texto.upper()
