"""Etapa 6, fase 3: BORRADOR TERLICA, la reconstruccion renglon por renglon.

Es la fase de mayor riesgo. La hoja trae una ESTRUCTURA del formulario
municipal -- renglones B1..B8 con su tarifa -- que NO es dato nuestro: es la
forma del formulario. El motor llena las columnas de datos (retencion,
concepto, actividad, NIT) en la fila cuya tarifa corresponde.

El peligro: un mes con mas grupos de una misma tarifa que filas disponibles.
Es el mismo problema que 3.2 resolvio en Python -- varios grupos por renglon
-- ahora expresado en filas de Excel. Si no cabe, tiene que DECIRLO, no
descartar el grupo en silencio.
"""

import warnings
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from motor_reteica.ingesta.sap_retenciones import RetencionERP
from motor_reteica.parametros.municipios.santa_marta import MUNICIPIO
from motor_reteica.pipeline import mapa_actividad, revisar
from motor_reteica.plantilla.deposito import depositar
from motor_reteica.reconstruccion import reconstruir
from motor_reteica.tipos import LineaAuxiliar

BASE = Path(__file__).parent / "fixtures" / "terlica_202607"


@pytest.fixture(scope="module")
def ctx():
    return revisar(BASE, nit="819002433", periodo="2026-07", municipio=MUNICIPIO)


@pytest.fixture(scope="module")
def papel(ctx, tmp_path_factory):
    salida = tmp_path_factory.mktemp("p") / "papel.xlsx"
    depositar(ctx, salida)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(salida)


def _ctx_con(ctx, lineas, erp):
    """Un contexto con OTRO mes de verdad.

    replace() a secas no basta: reconstruccion es un campo derivado y
    seguiria trayendo los grupos de julio. Hay que reconstruir.
    """
    mapa = mapa_actividad(ctx.borrador, erp, lineas, MUNICIPIO)
    return replace(ctx, lineas=lineas, filas_erp=erp,
                   reconstruccion=reconstruir(lineas, erp, MUNICIPIO, mapa))


def _filas_con_retencion(hoja):
    return {hoja.cell(row=f, column=8).value: f
            for f in range(10, 28)
            if hoja.cell(row=f, column=8).value}


def test_la_estructura_del_formulario_no_se_toca(papel):
    """B1..B8 y sus tarifas son la forma del formulario municipal, no dato
    nuestro. Reescribirlas seria inventar la estructura del impuesto."""
    hoja = papel["BORRADOR TERLICA"]
    assert hoja["B10"].value == "B1"
    assert hoja["C10"].value == 2
    assert hoja["B20"].value == "B7"
    assert hoja["C20"].value == 10


def test_cada_grupo_cae_en_una_fila_de_su_propia_tarifa(papel, ctx):
    hoja = papel["BORRADOR TERLICA"]
    for retencion, fila in _filas_con_retencion(hoja).items():
        por_mil = hoja.cell(row=fila, column=3).value
        grupo = next((g for g in ctx.reconstruccion.por_grupo.values()
                      if int(g.retencion_contable) == retencion), None)
        assert grupo is not None, "retencion %s no viene de ningun grupo" % retencion
        assert float(grupo.tarifa) * 1000 == por_mil, \
            "grupo al %s puesto en una fila de %s por mil" % (grupo.tarifa, por_mil)


def test_las_cinco_retenciones_del_mes_estan(papel):
    hoja = papel["BORRADOR TERLICA"]
    assert sorted(_filas_con_retencion(hoja)) == [1318, 11156, 35000, 37430, 389657]


def test_el_total_suma_lo_mismo_que_el_auxiliar(papel):
    hoja = papel["BORRADOR TERLICA"]
    assert sum(_filas_con_retencion(hoja)) == 474561


def test_las_formulas_de_la_hoja_siguen_vivas(papel):
    """D, E y G se recalculan solas a partir de H."""
    hoja = papel["BORRADOR TERLICA"]
    fila = min(_filas_con_retencion(hoja).values())
    assert str(hoja.cell(row=fila, column=4).value).startswith("=")
    assert str(hoja.cell(row=fila, column=7).value).startswith("=")


def test_no_queda_residuo_de_un_mes_anterior(ctx, tmp_path):
    """La plantilla trae las retenciones de julio en H16..H25."""
    lineas = [LineaAuxiliar(
        cuenta="2368010010", nit="900000001", tercero="UNICO",
        fecha_documento=date(2026, 8, 1), fecha_contabilizacion=date(2026, 8, 1),
        referencia="AG1", documento="1", concepto="COMPRA DE AGOSTO",
        retencion=Decimal("5000"))]
    corto = _ctx_con(ctx, lineas,
                     [RetencionERP(nit="900000001", codigo_ret="23",
                                   base=Decimal("500000"),
                                   retencion=Decimal("5000"))])
    salida = tmp_path / "agosto.xlsx"
    depositar(corto, salida)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        hoja = openpyxl.load_workbook(salida)["BORRADOR TERLICA"]

    assert sorted(_filas_con_retencion(hoja)) == [5000], \
        "quedaron retenciones de julio en el papel de agosto"


def test_un_grupo_que_no_cabe_se_declara_en_vez_de_perderse(ctx, tmp_path):
    """Mas grupos al 10 por mil que filas disponibles para esa tarifa.

    Es el caso de 3.2 en version Excel. Descartar el grupo en silencio
    produciria un papel que cuadra de menos sin decir por que.
    """
    lineas, erp = [], []
    for indice in range(12):   # mas que las filas al 10 por mil que hay
        nit = "9%08d" % indice
        lineas.append(LineaAuxiliar(
            cuenta="2368010010", nit=nit, tercero="P%s" % nit,
            fecha_documento=date(2026, 8, 1),
            fecha_contabilizacion=date(2026, 8, 1),
            referencia="AG%d" % indice, documento=str(indice),
            concepto="COMPRA", retencion=Decimal("1000")))
        erp.append(RetencionERP(nit=nit, codigo_ret="23",
                                base=Decimal("100000"), retencion=Decimal("1000")))
    muchos = _ctx_con(ctx, lineas, erp)

    salida = tmp_path / "muchos.xlsx"
    depositar(muchos, salida)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        hoja = openpyxl.load_workbook(salida)["BORRADOR TERLICA"]

    texto = " ".join(str(c.value) for fila in hoja.iter_rows()
                     for c in fila if c.value is not None).upper()
    assert "NO CUPO" in texto or "NO CABEN" in texto, \
        "los grupos que no caben deben quedar declarados en la hoja"
