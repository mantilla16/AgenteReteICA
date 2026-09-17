"""La interfaz web tiene que entregar el papel de la firma, no otro libro.

La etapa 6 cambio el entregable: el resultado es la plantilla real de la
firma, con sus nueve hojas y su logo. La CLI se cambio (--formato plantilla,
por defecto) y la API se quedo llamando a generar_papel(), el libro que se
inventaba el motor.

Nadie lo noto porque la API no tenia una sola prueba de extremo a extremo.
El error salio compilando el .exe y probandolo a mano: la interfaz --que es
lo que ve el usuario que hace doble clic-- devolvia el formato viejo.

Por eso esta prueba mira las HOJAS del archivo descargado y no que la
llamada no reviente: lo que fallaba devolvia 200 y un .xlsx perfectamente
valido, solo que el equivocado.
"""

import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from motor_reteica.api import servidor
from motor_reteica.api.servidor import app

BASE = Path(__file__).parent / "fixtures" / "terlica_202607"

# Un auditor con sesion valida. La API exige sesion desde que se le puso
# login, y esta prueba entra por el mismo camino que un auditor de verdad:
# con su cookie, pasando por el middleware. Lo unico doblado es la base --
# ninguna prueba de este proyecto habla con Postgres ni con la red.
USUARIO = {"id": "11111111-1111-1111-1111-111111111111",
           "usuario": "auditor", "nombre": "Auditor de prueba",
           "correo": "auditor@rbcol.co", "rol": "AUDITOR", "activo": True}

# Las de la plantilla de la firma. El libro viejo traia otras
# ("Caratula", "Controles", "Recalculo"...), asi que basta con exigir estas.
HOJAS_DE_LA_FIRMA = {"Check List", "DECLARACION", "REVISION ICA",
                     "BORRADOR TERLICA", "BALANCE", "AUX FISCAL",
                     "CUADRO RETEICA"}


@pytest.fixture(scope="module")
def corrida(tmp_path_factory):
    """Sube las fixtures por HTTP y devuelve lo que respondio la API.

    Devuelve el .xlsx descargado, el cuerpo del POST y lo que se mando a
    registrar, para que las pruebas de abajo miren cada cosa por separado.
    """
    papeles = tmp_path_factory.mktemp("papeles")
    registradas = []

    with pytest.MonkeyPatch.context() as mp:
        # Sin llave ni modelo local: la Revision Inteligente no entra, y la
        # prueba no depende de que la maquina que la corre tenga una llave.
        for var in ("ANTHROPIC_API_KEY", "OLLAMA_MODELO", "OLLAMA_URL",
                    "RETEICA_IA_PROVEEDOR"):
            mp.delenv(var, raising=False)

        mp.setattr(servidor, "_PAPELES", papeles)
        mp.setattr(servidor.db, "abrir", lambda: None)
        mp.setattr(servidor.db, "cerrar", lambda: None)
        mp.setattr(servidor.db, "usuario_de_sesion", lambda _hash: USUARIO)
        mp.setattr(servidor.db, "guardar_revision",
                   lambda **campos: registradas.append(campos))
        mp.setattr(servidor.db, "revision_de",
                   lambda corrida, usuario_id: {
                       "corrida": corrida, "usuario_id": usuario_id,
                       "nit": "819002433", "periodo": "2026-07"})

        with TestClient(app) as cliente:
            cliente.cookies.set(servidor.COOKIE, "token-de-prueba")
            archivos = [("archivos", (r.name, r.read_bytes()))
                        for r in sorted(BASE.iterdir()) if r.is_file()]
            archivos += [("archivos", (r.name, r.read_bytes()))
                         for r in sorted((BASE / "facturas").iterdir())]
            respuesta = cliente.post(
                "/analizar", files=archivos,
                data={"nit": "819002433", "periodo": "2026-07",
                      "municipio": "santa_marta"})
            assert respuesta.status_code == 200, respuesta.text
            cuerpo = respuesta.json()

            papel = cliente.get(cuerpo["descarga"])
            assert papel.status_code == 200, papel.text

    destino = tmp_path_factory.mktemp("api") / "papel.xlsx"
    destino.write_bytes(papel.content)
    return {"papel": destino, "cuerpo": cuerpo, "registradas": registradas,
            "deposito": papeles, "respuesta_descarga": papel}


@pytest.fixture(scope="module")
def descargado(corrida):
    return corrida["papel"]


def test_la_api_devuelve_el_papel_de_la_firma(descargado):
    libro = openpyxl.load_workbook(descargado)
    faltan = HOJAS_DE_LA_FIRMA - set(libro.sheetnames)
    assert not faltan, "la API no entrego la plantilla; faltan %s" % faltan


def test_el_papel_de_la_api_conserva_el_logo(descargado):
    """fidelidad.py restituye el EMF; openpyxl solo no lo lograria."""
    partes = set(zipfile.ZipFile(descargado).namelist())
    assert "xl/media/image1.emf" in partes


def test_el_papel_de_la_api_trae_las_cifras_del_motor(descargado):
    """D11, version final: lo que cruza de hoja va como numero.

    Antes esta prueba exigia que F20 fuera un SUMIF. Con formula, lo unico
    verificable era la cadena; el numero podia salir en blanco o en #N/D y
    la prueba seguia en verde. Ahora se comprueba la cifra.
    """
    hoja = openpyxl.load_workbook(descargado)["REVISION ICA"]
    assert isinstance(hoja["F20"].value, (int, float)), hoja["F20"].value
    assert hoja["N12"].value + hoja["N13"].value == 474561


# --------------------------------------------------------------------------
# Lo que la corrida deja atras: registro, archivo y lo que se descarta
# --------------------------------------------------------------------------

def test_la_revision_queda_registrada_en_el_historial(corrida):
    """Se hacen varias revisiones y hay que saber cuales se hicieron."""
    assert len(corrida["registradas"]) == 1
    fila = corrida["registradas"][0]
    assert fila["nit"] == "819002433"
    assert fila["periodo"] == "2026-07"
    assert fila["usuario_id"] == USUARIO["id"]
    assert fila["corrida"] == corrida["cuerpo"]["corrida"]
    assert corrida["cuerpo"]["registrada"] is True


def test_el_registro_guarda_el_resumen_completo(corrida):
    """La fila tiene que bastar para repintar la revision sin rehacerla."""
    resumen = corrida["registradas"][0]["resumen"]
    assert resumen["controles"], "sin controles no se puede repintar nada"
    assert "semaforo" in resumen and "conclusion" in resumen


def test_el_papel_queda_en_el_deposito_durable(corrida):
    """No en /tmp: ahi el sistema lo borra y el historial queda apuntando
    a un archivo que ya no esta."""
    esperado = corrida["deposito"] / ("%s.xlsx" % corrida["cuerpo"]["corrida"])
    assert esperado.exists()


def test_los_documentos_del_cliente_no_se_quedan_en_el_servidor(corrida):
    """Se guarda el papel, no la evidencia cruda. Fue una decision explicita:
    son cifras de un cliente y no hay razon para acumularlas."""
    trabajo = servidor._RAIZ / corrida["cuerpo"]["corrida"]
    assert not trabajo.exists()


def test_la_descarga_llega_con_nombre_archivable(corrida):
    """El auditor guarda esto en su carpeta: el id de la corrida no le sirve."""
    disposicion = corrida["respuesta_descarga"].headers.get("content-disposition", "")
    assert "819002433" in disposicion and "202607" in disposicion


def test_sin_sesion_la_api_no_analiza():
    """La regresion que rompio este archivo: al poner login, /analizar quedo
    detras del middleware y nadie lo habia comprobado."""
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(servidor.db, "abrir", lambda: None)
        mp.setattr(servidor.db, "cerrar", lambda: None)
        mp.setattr(servidor.db, "usuario_de_sesion", lambda _hash: None)
        with TestClient(app) as cliente:
            respuesta = cliente.post(
                "/analizar",
                files=[("archivos", ("x.pdf", b"%PDF-1.4 no importa"))],
                data={"nit": "819002433", "periodo": "2026-07"})
    assert respuesta.status_code != 200, "sin sesion no se revisa nada"
