"""Etapa 6, fase 1: depositar los datos crudos en la plantilla.

Tres hojas que son FUENTE, no resultado: no dependen de formulas de otras
hojas, asi que son las de menor riesgo para empezar.

El peligro real de esta fase no es escribir mal: es NO BORRAR. La plantilla
viene con los 12 renglones de julio. Un mes con 3 lineas que solo sobreescriba
las 3 primeras deja las 9 restantes de julio en el papel de agosto, con su
pinta de dato bueno. Ese es el fallo que estas pruebas persiguen.
"""

import warnings
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from motor_reteica.ingesta.sap_retenciones import RetencionERP
from motor_reteica.parametros.municipios.santa_marta import MUNICIPIO
from motor_reteica.pipeline import revisar
from motor_reteica.plantilla.deposito import depositar
from motor_reteica.tipos import LineaAuxiliar

BASE = Path(__file__).parent / "fixtures" / "terlica_202607"


@pytest.fixture(scope="module")
def ctx():
    return revisar(BASE, nit="819002433", periodo="2026-07", municipio=MUNICIPIO)


@pytest.fixture(scope="module")
def papel(ctx, tmp_path_factory):
    salida = tmp_path_factory.mktemp("plantilla") / "papel.xlsx"
    depositar(ctx, salida)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(salida)


# --------------------------------------------------------------------------
# AUX FISCAL
# --------------------------------------------------------------------------

def test_aux_fiscal_recibe_las_doce_lineas(papel):
    hoja = papel["AUX FISCAL"]
    cuentas = [hoja.cell(row=f, column=2).value for f in range(7, 19)]
    assert all(str(c).startswith("2368") for c in cuentas), cuentas


def test_aux_fiscal_conserva_el_signo_credito_del_erp(papel):
    """El auxiliar de SAP trae los importes en negativo. El motor los guarda
    en positivo; al depositarlos hay que devolverles el signo o el papel deja
    de parecerse al documento del cliente."""
    hoja = papel["AUX FISCAL"]
    importes = [hoja.cell(row=f, column=9).value for f in range(7, 19)]
    assert all(i < 0 for i in importes), importes
    assert sum(importes) == -474561


def test_aux_fiscal_lleva_concepto_y_referencia(papel):
    hoja = papel["AUX FISCAL"]
    textos = [hoja.cell(row=f, column=13).value for f in range(7, 19)]
    referencias = [str(hoja.cell(row=f, column=8).value) for f in range(7, 19)]
    assert 'REPARACION MANGUERA DE 6"' in textos
    assert "FE338057" in referencias


def test_el_total_del_auxiliar_sigue_siendo_una_formula(papel):
    """Congelarlo a valor mataria la hoja: es un papel vivo (D11)."""
    hoja = papel["AUX FISCAL"]
    assert str(hoja["I19"].value).upper().startswith("=SUM")
    assert hoja["B19"].value == "TOTAL"


# --------------------------------------------------------------------------
# CUADRO RETEICA
# --------------------------------------------------------------------------

def test_cuadro_reteica_recibe_los_cinco_terceros(papel):
    hoja = papel["CUADRO RETEICA"]
    nits = [str(hoja.cell(row=f, column=2).value) for f in range(5, 10)]
    assert sorted(nits) == ["800193573", "811033997", "830028245",
                            "900392924", "901670478"]


def test_cuadro_reteica_lleva_base_y_retencion(papel):
    hoja = papel["CUADRO RETEICA"]
    bases = [hoja.cell(row=f, column=7).value for f in range(5, 10)]
    retenciones = [hoja.cell(row=f, column=8).value for f in range(5, 10)]
    assert sum(bases) == 49012569
    assert sum(retenciones) == 474561


# --------------------------------------------------------------------------
# BALANCE
# --------------------------------------------------------------------------

def test_balance_recibe_las_cuentas_2368(papel):
    hoja = papel["BALANCE"]
    filas = {}
    for f in range(5, hoja.max_row + 1):
        cuenta = hoja.cell(row=f, column=3).value
        if cuenta:
            filas[str(cuenta)] = hoja.cell(row=f, column=9).value
    assert filas.get("2368010007") == 36318
    assert filas.get("2368010010") == 438243


def test_el_balance_distingue_evidencia_de_lo_verificado(papel):
    """Cambio de forma, no de exigencia.

    Antes la hoja traia solo las 2368 y se rotulaba EXTRACTO, porque
    depositar 349 filas que el motor nunca miro daria a entender que las
    reviso. Ahora se transcribe el balance COMPLETO --el papel de trabajo es
    la evidencia, y sin ella el lector no puede comprobar ningun saldo-- asi
    que la distincion tiene que hacerla el texto: se dice que solo las 2368
    se verifican, y la columna LECTURA DEL MOTOR lo marca cuenta por cuenta.
    """
    hoja = papel["BALANCE"]
    texto = " ".join(str(c.value) for fila in hoja.iter_rows(max_row=4)
                     for c in fila if c.value is not None).upper()
    assert "EVIDENCIA" in texto
    assert "SOLO VERIFICA LAS CUENTAS 2368" in texto
    assert "LECTURA DEL MOTOR" in texto


# --------------------------------------------------------------------------
# El fallo peligroso: residuo del mes anterior
# --------------------------------------------------------------------------

def _linea(nit, referencia, retencion):
    return LineaAuxiliar(
        cuenta="2368010010", nit=nit, tercero="PROVEEDOR " + nit,
        fecha_documento=date(2026, 8, 3), fecha_contabilizacion=date(2026, 8, 3),
        referencia=referencia, documento="9", concepto="COMPRA DE AGOSTO",
        retencion=Decimal(retencion))


def test_un_mes_mas_corto_no_deja_residuo_del_anterior(ctx, tmp_path):
    """El caso que de verdad importa: agosto con 3 lineas sobre una plantilla
    que trae las 12 de julio."""
    from dataclasses import replace
    lineas = [_linea("900000001", "AG1", "1000"),
              _linea("900000002", "AG2", "2000"),
              _linea("900000003", "AG3", "3000")]
    erp = [RetencionERP(nit=l.nit, codigo_ret="23",
                        base=l.retencion * 100, retencion=l.retencion)
           for l in lineas]
    corto = replace(ctx, lineas=lineas, filas_erp=erp)

    salida = tmp_path / "agosto.xlsx"
    depositar(corto, salida)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        hoja = openpyxl.load_workbook(salida)["AUX FISCAL"]

    referencias = [hoja.cell(row=f, column=8).value for f in range(7, 19)]
    assert referencias[:3] == ["AG1", "AG2", "AG3"]
    assert all(r is None for r in referencias[3:]), \
        "quedo residuo de julio en el papel de agosto: %s" % referencias[3:]


def test_el_total_se_reajusta_al_numero_de_lineas(ctx, tmp_path):
    from dataclasses import replace
    lineas = [_linea("900000001", "AG1", "1000")]
    corto = replace(ctx, lineas=lineas,
                    filas_erp=[RetencionERP(nit="900000001", codigo_ret="23",
                                            base=Decimal("100000"),
                                            retencion=Decimal("1000"))])
    salida = tmp_path / "agosto.xlsx"
    depositar(corto, salida)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        hoja = openpyxl.load_workbook(salida)["AUX FISCAL"]
    # una sola linea: el total va en la fila siguiente y suma solo esa
    assert hoja["B8"].value == "TOTAL"
    assert hoja["I8"].value == "=SUM(I7:I7)"
    assert hoja["B19"].value is None, "quedo el TOTAL viejo de julio"


def test_el_logo_sobrevive_al_deposito(ctx, tmp_path):
    """La fase 1 no puede deshacer lo que resolvio fidelidad.py."""
    import zipfile
    salida = tmp_path / "papel.xlsx"
    depositar(ctx, salida)
    with zipfile.ZipFile(salida) as z:
        assert "xl/media/image1.emf" in z.namelist()


# --------------------------------------------------------------------------
# Hallazgos del loop de validacion (vuelta 1)
# --------------------------------------------------------------------------

def test_el_texto_de_la_cuenta_muestra_el_porcentaje_no_los_digitos(papel):
    """V1: se escribia 'Impuest ICA Reten 0007' -- los ultimos digitos de la
    cuenta -- donde SAP y la plantilla dicen 'Impuest ICA Reten 7%'."""
    hoja = papel["AUX FISCAL"]
    textos = {hoja.cell(row=f, column=3).value for f in range(7, 19)
              if hoja.cell(row=f, column=3).value}
    assert textos == {"Impuest ICA Reten 7%", "Impuest ICA Reten 10%"}, textos


def test_la_nota_del_motor_dice_la_tarifa_de_lo_que_cruzo(papel):
    """La tarifa sigue estando, ahora en la columna del motor.

    Y solo sobre las cuentas que entraron al cruce: una tarifa junto a una
    cuenta que no se verifico daria a entender que si. Esa exigencia no
    cambio, cambio el lugar -- la columna 4 ahora trae el texto del cliente,
    porque la fuente se transcribe intacta.
    """
    hoja = papel["BALANCE"]
    cruzadas = [hoja.cell(row=f, column=11).value
                for f in range(5, hoja.max_row + 1)
                if str(hoja.cell(row=f, column=11).value or "").startswith("CRUZADA")]
    assert cruzadas, "no quedo ninguna cuenta marcada como cruzada"
    assert all("%" in t for t in cruzadas), cruzadas


def test_el_nit_de_la_caratula_lleva_digito_de_verificacion(papel):
    """V2: el papel de la firma identifica al cliente como 819002433-6.
    Escribir el NIT sin su DV lo deja incompleto en un documento que se firma."""
    assert str(papel["Check List"]["D3"].value) == "819002433-6"
