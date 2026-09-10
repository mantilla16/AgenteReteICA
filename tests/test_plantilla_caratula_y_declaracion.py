"""Etapa 6, fases 2 y 4: caratula, checklist, declaracion y facturas.

La marca del checklist va en fuente Webdings, donde 'a' se dibuja como un
visto y 'x' como una equis. Se conservan esas letras.

El punto de auditoria de este bloque: el motor llena ELABORADO POR porque el
papel lo prepara el, pero NUNCA llena REVISADO POR. Escribir ahi un nombre
seria fabricar evidencia de una revision que no ocurrio.
"""

import warnings
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from motor_reteica.parametros.municipios.santa_marta import MUNICIPIO
from motor_reteica.pipeline import revisar
from motor_reteica.plantilla.deposito import depositar

BASE = Path(__file__).parent / "fixtures" / "terlica_202607"


def _solo_fecha(valor):
    """openpyxl devuelve datetime; lo que importa es el dia."""
    return valor.date() if hasattr(valor, "date") else valor


@pytest.fixture(scope="module")
def papel(tmp_path_factory):
    ctx = revisar(BASE, nit="819002433", periodo="2026-07", municipio=MUNICIPIO)
    salida = tmp_path_factory.mktemp("p") / "papel.xlsx"
    depositar(ctx, salida)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(salida)


# --------------------------------------------------------------------------
# Check List: caratula
# --------------------------------------------------------------------------

def test_la_caratula_identifica_al_cliente(papel):
    hoja = papel["Check List"]
    assert "819002433" in str(hoja["D3"].value)
    assert "GRANELES" in str(hoja["D2"].value).upper()


def test_la_caratula_trae_periodo_y_vencimiento(papel):
    hoja = papel["Check List"]
    assert str(hoja["D7"].value).strip().lower() == "julio"
    assert _solo_fecha(hoja["D5"].value) == date(2026, 8, 14)


def test_el_motor_firma_como_elaborado_pero_no_como_revisado(papel):
    """Llenar REVISADO POR seria fabricar una revision que no ocurrio."""
    hoja = papel["Check List"]
    assert hoja["D10"].value, "ELABORADO POR deberia quedar diligenciado"
    assert not hoja["D11"].value, \
        "REVISADO POR debe quedar VACIO: nadie ha revisado este papel"
    assert not hoja["D12"].value, "REVISADO POR DIRECTOR tambien vacio"


# --------------------------------------------------------------------------
# Check List: marcas del checklist
# --------------------------------------------------------------------------

def test_los_insumos_obtenidos_quedan_marcados_con_visto(papel):
    hoja = papel["Check List"]
    assert hoja["E18"].value == "a", "balance obtenido -> visto"
    assert hoja["E19"].value == "a", "borrador obtenido -> visto"
    assert hoja["E22"].value == "a", "auxiliar 2368 obtenido -> visto"


def test_los_insumos_no_obtenidos_quedan_marcados_con_equis(papel):
    hoja = papel["Check List"]
    assert hoja["E24"].value == "x", "el pago del mes anterior no se obtuvo"


def test_lo_que_esta_fuera_de_alcance_se_marca_na_y_no_equis(papel):
    """El auxiliar de ingresos es de AUTORRETENCION: otro papel, otro riesgo
    (spec 2.2). Marcarlo con equis diria que falto algo que nunca se pidio."""
    assert papel["Check List"]["E25"].value == "N/A"


def test_la_marca_conserva_la_fuente_webdings(papel):
    """Si se pierde la fuente, el visto se lee como la letra 'a'."""
    assert papel["Check List"]["E18"].font.name == "Webdings"


# --------------------------------------------------------------------------
# DECLARACION
# --------------------------------------------------------------------------

def test_la_declaracion_recibe_los_cinco_renglones(papel):
    hoja = papel["DECLARACION"]
    actividades = [str(hoja.cell(row=f, column=9).value) for f in range(3, 8)]
    assert sorted(actividades) == ["4669", "5224", "7490", "9609", "9903"]


def test_la_declaracion_cuadra_con_el_borrador(papel):
    hoja = papel["DECLARACION"]
    bases = [hoja.cell(row=f, column=11).value for f in range(3, 8)]
    impuestos = [hoja.cell(row=f, column=12).value for f in range(3, 8)]
    assert sum(bases) == 49013000
    assert sum(impuestos) == 474000


def test_los_totales_de_la_declaracion_siguen_siendo_formula(papel):
    hoja = papel["DECLARACION"]
    assert str(hoja["K8"].value).upper().startswith("=SUM")
    assert str(hoja["L8"].value).upper().startswith("=SUM")


# --------------------------------------------------------------------------
# Validacion de facturas  (es C11 hecho papel)
# --------------------------------------------------------------------------

def test_las_facturas_cotejadas_llegan_al_papel(papel):
    hoja = papel["Validación de facturas"]
    facturas = [str(hoja.cell(row=f, column=3).value) for f in range(27, 30)]
    assert "FE10" in facturas
    assert "FE338057" in facturas


def test_la_factura_lleva_base_tarifa_y_valor_contable(papel):
    hoja = papel["Validación de facturas"]
    for fila in range(27, 30):
        if str(hoja.cell(row=fila, column=3).value) == "FE10":
            assert hoja.cell(row=fila, column=7).value == 5000000
            assert hoja.cell(row=fila, column=8).value == 0.007
            assert hoja.cell(row=fila, column=10).value == 35000
            return
    pytest.fail("no se deposito FE10")


def test_el_calculo_del_auditor_y_la_diferencia_siguen_siendo_formula(papel):
    """La columna 'Calculo de auditor' es el recalculo vivo del papel."""
    hoja = papel["Validación de facturas"]
    assert str(hoja["I27"].value).startswith("=")
    assert str(hoja["K27"].value).startswith("=")


def test_la_fecha_de_la_factura_es_la_del_pdf_no_la_del_auxiliar(papel):
    """FE10 se genero el 26/06 y el auxiliar la registra el 01/07. El papel
    debe mostrar la del documento, que es lo que hace visible el desfase."""
    hoja = papel["Validación de facturas"]
    for fila in range(27, 30):
        if str(hoja.cell(row=fila, column=3).value) == "FE10":
            leida = _solo_fecha(hoja.cell(row=fila, column=4).value)
            assert leida == date(2026, 6, 26)
            return
    pytest.fail("no se deposito FE10")
