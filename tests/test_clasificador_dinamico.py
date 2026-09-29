"""El clasificador y el lector reconocen un papel del cliente (formato Mov SAP).

Este formato rompia el motor: encabezado tras 5 filas vacias, la retencion en
una columna 'Retencion' (no 'Importe en ML'), filas de subtotal intercaladas
('Cuenta 2368010005 - ...') y la cuenta puesta solo en la primera fila de cada
bloque. Se reconstruye sintetico para correr en CI sin la carpeta del cliente.
"""

from decimal import Decimal

import openpyxl
import pytest

from motor_reteica.ingesta.auxiliar import leer_auxiliar

_HEADER = ["Cuenta", "Texto breve", "Asignación", "Tercero",
           "Fecha de documento", "Referencia", "Base de Retención",
           "Retención", "Período contable", "CIIU", "Nº documento", "Texto",
           "Sociedad"]

_FILAS = [
    # detalle: cuenta presente
    [2368010005, "Impuest ICA Reten 5%", 36560048, "RIBAUTT OSORIO GRACIELA",
     "27.08.2026", "FVLR1068", 2088400, 10442, 7, "5611-5313", 5100003125,
     "ALIMENTACION", "DA47"],
    # continuacion: MISMA cuenta, celda en blanco (se arrastra)
    [None, None, 36560048, "RIBAUTT OSORIO GRACIELA", "14.08.2026", "FVLR1057",
     2446800, 12234, 7, "5611-5313", 5100003112, "ALIMENTACION", None],
    # subtotal: se descarta (no doblar)
    ["Cuenta 2368010005 - Retención Impuest ICA Reten 5%", None, None, None,
     None, None, 4535200, 22676, None, None, None, None, None],
    # otro bloque
    [2368010007, "Impuest ICA Reten 7%", 8698821, "DELGADO OROZCO",
     "31.08.2026", "DE252", 3180000, 22260, 7, 7110, 5100003099, "HONORARIOS",
     "DA47"],
    ["Cuenta 2368010007 - Retención Impuest ICA Reten 7%", None, None, None,
     None, None, 3180000, 22260, None, None, None, None, None],
]


@pytest.fixture
def archivo(tmp_path):
    """Workbook con 5 filas vacias antes del encabezado y 3 hojas, como el real."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mov SAP"
    for _ in range(5):
        ws.append([])
    ws.append(_HEADER)
    for fila in _FILAS:
        ws.append(fila)
    wb.create_sheet("Bce 2368")
    wb.create_sheet("DECLARACION ")
    ruta = tmp_path / "Copia de Mov SAP 2368 Agosto 2026.xlsx"
    wb.save(ruta)
    return ruta


def test_el_lector_arrastra_cuenta_y_descarta_subtotales(archivo):
    lineas = leer_auxiliar(archivo)
    # Tres transacciones reales (no los subtotales), con la cuenta arrastrada.
    assert len(lineas) == 3
    total = sum((l.retencion for l in lineas), Decimal("0"))
    assert total == Decimal("44936")          # 10442 + 12234 + 22260
    # La continuacion heredo la cuenta 2368010005.
    cuentas = [l.cuenta for l in lineas]
    assert cuentas == ["2368010005", "2368010005", "2368010007"]


def test_el_clasificador_reconoce_el_papel(archivo):
    # servidor importa fastapi/psycopg; si no estan, se omite esta parte.
    pytest.importorskip("fastapi")
    pytest.importorskip("psycopg")
    from motor_reteica.api.servidor import _rol_xlsx
    assert _rol_xlsx(archivo) == "auxiliar"
