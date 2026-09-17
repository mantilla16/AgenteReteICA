"""El cruce universal: total declarado vs total auxiliar 2368.

Es lo que Robinson hace a mano al final y lo que funciona en cualquier
municipio -- solo depende de dos numeros. No sustituye a los cruces por
renglon; los complementa: cuando el motor no reconoce la estructura del
borrador (San Alberto, cualquier municipio nuevo), este cruce es lo unico
que sostiene la conclusion.

Es un dato del auditor: la cifra viene de la tarjeta de confirmacion en el
tablero, no de la extraccion automatica. D9 -- el motor propone, el auditor
confirma.
"""

import warnings
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from motor_reteica.parametros.municipios.santa_marta import MUNICIPIO
from motor_reteica.pipeline import revisar
from motor_reteica.plantilla.deposito import depositar

BASE = Path(__file__).parent / "fixtures" / "terlica_202607"


def _generar(tmp_path, total_confirmado):
    ctx = revisar(BASE, nit="819002433", periodo="2026-07", municipio=MUNICIPIO)
    salida = tmp_path / "papel.xlsx"
    depositar(ctx, salida, total_declarado_confirmado=total_confirmado)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(salida)["REVISION ICA"]


def _bloque_cruce(hoja) -> dict:
    """Devuelve {etiqueta: valor} para las filas del cruce."""
    for fila in range(28, hoja.max_row + 5):
        cabecera = hoja.cell(row=fila, column=2).value
        if cabecera == "CRUCE UNIVERSAL":
            return {
                "declarado": hoja.cell(row=fila + 1, column=4).value,
                "auxiliar":  hoja.cell(row=fila + 2, column=4).value,
                "diferencia": hoja.cell(row=fila + 3, column=4).value,
                "veredicto": hoja.cell(row=fila + 3, column=5).value,
            }
    return {}


def test_sin_total_confirmado_no_hay_cruce(tmp_path):
    """Sin la cifra del auditor no se afirma que cuadre. D9."""
    ctx = revisar(BASE, nit="819002433", periodo="2026-07", municipio=MUNICIPIO)
    salida = tmp_path / "papel.xlsx"
    depositar(ctx, salida)  # sin total_declarado_confirmado
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        hoja = openpyxl.load_workbook(salida)["REVISION ICA"]
    assert _bloque_cruce(hoja) == {}


def test_con_total_confirmado_el_papel_muestra_el_cruce(tmp_path):
    """El bloque queda escrito al final de REVISION ICA."""
    hoja = _generar(tmp_path, Decimal("474000"))
    bloque = _bloque_cruce(hoja)
    assert bloque["declarado"] == 474000
    assert bloque["auxiliar"] == 474561    # suma real del auxiliar de TERLICA
    assert bloque["diferencia"] == -561    # 474000 - 474561


def test_diferencia_pequena_se_lee_como_cuadre(tmp_path):
    """Los formularios redondean a miles -- una diferencia de -561 pesos NO
    es un hallazgo, es el redondeo. El papel lo dice."""
    hoja = _generar(tmp_path, Decimal("474000"))
    assert "CUADRA" in _bloque_cruce(hoja)["veredicto"]


def test_diferencia_grande_se_marca_para_revisar(tmp_path):
    """Si el declarado se aleja del auxiliar mas de 1.000 pesos, el papel
    no lo llama CUADRA -- lo marca."""
    hoja = _generar(tmp_path, Decimal("400000"))
    veredicto = _bloque_cruce(hoja)["veredicto"]
    assert "NO CUADRA" in veredicto
