import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from motor_reteica.papel_excel import generar_papel
from motor_reteica.parametros.municipios.santa_marta import MUNICIPIO
from motor_reteica.pipeline import revisar

BASE = Path(__file__).parent / "fixtures" / "terlica_202607"


@pytest.fixture(scope="module")
def ctx():
    return revisar(BASE, nit="819002433", periodo="2026-07", municipio=MUNICIPIO)


@pytest.fixture(scope="module")
def libro(ctx, tmp_path_factory):
    salida = tmp_path_factory.mktemp("papel") / "papel.xlsx"
    return openpyxl.load_workbook(generar_papel(salida, ctx))


def _texto(hoja):
    return " ".join(str(c.value) for fila in hoja.iter_rows() for c in fila
                    if c.value is not None)


def test_genera_las_hojas_del_papel(libro):
    assert {"Notas", "Caratula", "Controles", "Recalculo", "Cruces",
            "Excepciones", "Parametros"} == set(libro.sheetnames)


def test_notas_es_la_primera_hoja(libro):
    """5.4: las conclusiones largas no se leen. La puerta de entrada va
    primero, con el semaforo arriba."""
    assert libro.sheetnames[0] == "Notas"
    texto = _texto(libro["Notas"])
    assert "SOLIDEZ DE LA REVISION" in texto
    assert any(color in texto for color in ("ROJO", "AMARILLO", "VERDE"))


def test_notas_lista_los_controles_no_ejecutados(libro):
    """Un control sin ejecutar tiene que verse en la primera hoja, no solo
    enterrado en la conclusion."""
    texto = _texto(libro["Notas"])
    assert "NO EJECUTADO" in texto
    assert "C7" in texto


def test_la_caratula_identifica_cliente_periodo_y_municipio(libro):
    texto = _texto(libro["Caratula"])
    assert "819002433" in texto
    assert "2026-07" in texto
    assert "Santa Marta" in texto


def test_la_conclusion_nombra_los_controles_no_ejecutados(libro):
    texto = _texto(libro["Caratula"])
    assert "C7" in texto and "C12" in texto
    assert "NO es concluyente" in texto


def test_la_conclusion_no_afirma_ausencia_de_diferencias(libro):
    """Con controles sin ejecutar no puede concluirse limpio."""
    assert "no presenta diferencias" not in _texto(libro["Caratula"])


def test_los_controles_traen_los_quince_codigos(libro):
    texto = _texto(libro["Controles"])
    for numero in range(0, 15):
        assert "C%d" % numero in texto


def test_el_recalculo_reproduce_las_anclas(libro):
    texto = _texto(libro["Recalculo"])
    assert "474561" in texto.replace(".0", "")
    assert "474000" in texto.replace(".0", "")


def test_parametros_declara_el_estado_de_las_tarifas(libro):
    assert "PENDIENTE_VALIDACION_ESTATUTO" in _texto(libro["Parametros"])


def test_parametros_registra_las_huellas(libro):
    texto = _texto(libro["Parametros"])
    assert "SHA256" in texto.upper()


def test_parametros_marca_los_insumos_no_obtenidos(libro):
    assert "NO OBTENIDO" in _texto(libro["Parametros"])


def test_el_papel_declara_la_limitacion_de_integridad(libro):
    texto = " ".join(_texto(libro[nombre]) for nombre in libro.sheetnames)
    assert "integridad" in texto.lower()


def test_las_excepciones_salen_ordenadas_por_severidad(libro):
    texto = _texto(libro["Excepciones"])
    assert "OBSERVACION" in texto or "AVISO" in texto


def test_generar_papel_no_revienta_con_un_control_atestado(tmp_path):
    """Bug real encontrado al construir la API web: _RELLENO_ESTADO no tenia
    entrada para Estado.ATESTADO y hoja_controles reventaba con KeyError la
    primera vez que alguien atestara una tarifa -- nunca se habia ejercitado
    porque nadie ha atestado en ninguna corrida real todavia."""
    destino = tmp_path / "con_atestacion"
    shutil.copytree(BASE, destino)
    tarifas = [{"municipio": "Santa Marta", "actividad": c, "tarifa": t,
               "vigencia_desde": "2026-01-01", "acuerdo": "Acuerdo 013 de 2024",
               "articulo": "52"}
              for c, t in {"9609": "0.007", "7490": "0.007", "4669": "0.010",
                          "5224": "0.010", "9903": "0.010"}.items()]
    (destino / "atestacion.json").write_text(json.dumps({
        "declarada_por": "analitica@rbcol.co", "fecha": "2026-09-09",
        "tarifas": tarifas, "cuentas": []}, ensure_ascii=False), encoding="utf-8")

    ctx = revisar(destino, nit="819002433", periodo="2026-07", municipio=MUNICIPIO)
    c7 = next(r for r in ctx.resultados if r.codigo == "C7")
    assert c7.estado.value == "ATESTADO"

    salida = generar_papel(destino / "papel.xlsx", ctx)  # no debe lanzar KeyError
    texto = " ".join(str(c.value) for fila in openpyxl.load_workbook(salida)["Controles"]
                     .iter_rows() for c in fila if c.value is not None)
    assert "ATESTADO" in texto


def test_excepciones_declara_el_tipo_de_referencia(libro):
    """M13: la columna RENGLON ya no es ambigua sin decir de que tipo es."""
    texto = _texto(libro["Excepciones"])
    assert "TIPO DE REF." in texto.upper()
    assert "Renglon" in texto or "Cuenta" in texto or "NIT" in texto


def test_parametros_declara_que_no_hubo_manifiesto_para_esta_corrida(libro):
    """1.4.b: la fixture corre en modo legado (sin manifiesto.json)."""
    texto = _texto(libro["Parametros"])
    assert "MANIFIESTO" in texto.upper()
    assert "Sin manifiesto.json" in texto
