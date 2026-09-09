"""Generador del papel de trabajo en formato AR-FO.

La conclusion no se redacta aqui: viene de hallazgos.consolidar, derivada del
estado de los controles. Este modulo solo la imprime.
"""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import motor_reteica
from motor_reteica.hallazgos import LIMITACION_INTEGRIDAD
from motor_reteica.semaforo import evaluar
from motor_reteica.tipos import Estado, Severidad
from motor_reteica.version import huella_del_codigo

AZUL = "001871"
GRIS = "F2F2F2"
VERDE = "C6EFCE"
ROJO = "FFC7CE"
AMBAR = "FFEB9C"

_RELLENO_ESTADO = {
    Estado.OK: VERDE,
    Estado.FALLA: ROJO,
    Estado.NO_EJECUTADO: AMBAR,
}

_BORDE = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def _titulo(hoja, fila, texto, ancho=8):
    celda = hoja.cell(row=fila, column=1, value=texto)
    celda.font = Font(name="Lato", size=12, bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=AZUL)
    hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
    return fila + 1


def _encabezados(hoja, fila, etiquetas):
    for columna, etiqueta in enumerate(etiquetas, start=1):
        celda = hoja.cell(row=fila, column=columna, value=etiqueta)
        celda.font = Font(name="Lato", size=10, bold=True)
        celda.fill = PatternFill("solid", fgColor=GRIS)
        celda.border = _BORDE
    return fila + 1


def _fila(hoja, fila, valores, relleno=None):
    for columna, valor in enumerate(valores, start=1):
        celda = hoja.cell(row=fila, column=columna, value=valor)
        celda.font = Font(name="Lato", size=10)
        celda.border = _BORDE
        if relleno:
            celda.fill = PatternFill("solid", fgColor=relleno)
    return fila + 1


def _nota_alcance(hoja, fila):
    celda = hoja.cell(row=fila + 1, column=1, value=LIMITACION_INTEGRIDAD)
    celda.font = Font(name="Lato", size=9, italic=True)
    celda.alignment = Alignment(wrap_text=True, vertical="top")
    hoja.merge_cells(start_row=fila + 1, start_column=1,
                     end_row=fila + 3, end_column=8)


def _hoja_caratula(libro, ctx):
    hoja = libro.create_sheet("Caratula")
    hoja.column_dimensions["A"].width = 34
    for letra in "BCDEFGH":
        hoja.column_dimensions[letra].width = 18

    fila = _titulo(hoja, 1, "PAPEL DE TRABAJO - REVISION DE RETEICA")
    datos = [
        ("NOMBRE DEL CLIENTE", ctx.borrador.razon_social),
        ("NIT", ctx.nit),
        ("TIPO DE TRABAJO", "Revisoria fiscal"),
        ("NOMBRE DEL PAPEL DE TRABAJO", "Revision Reteica"),
        ("PERIODO DE REVISION", ctx.periodo),
        ("MUNICIPIO", ctx.municipio.nombre),
        ("NORMATIVIDAD APLICABLE", "NIIF Plenas"),
        ("CIFRAS EXPRESADAS EN", "Pesos colombianos"),
        ("NUMERO DE FORMULARIO", ctx.borrador.numero_formulario),
        ("FIRMA DE REVISOR FISCAL EN EL BORRADOR",
         "Si" if ctx.borrador.firma_revisor_fiscal else "No"),
        # M7: sin esto el papel no se puede reproducir. La huella lleva
        # '+sucio' si se genero sobre codigo sin commitear.
        ("VERSION DEL MOTOR", motor_reteica.__version__),
        ("HUELLA DEL CODIGO", huella_del_codigo()),
    ]
    for etiqueta, valor in datos:
        fila = _fila(hoja, fila, [etiqueta, valor])

    fila += 1
    fila = _titulo(hoja, fila, "OBJETIVO")
    fila = _fila(hoja, fila, [
        "Verificar la razonabilidad del borrador de la declaracion mensual de "
        "ReteICA preparado por la compania, mediante la reconstruccion "
        "independiente de la liquidacion a partir de los libros."])

    fila += 1
    fila = _titulo(hoja, fila, "CONCLUSION")
    celda = hoja.cell(row=fila, column=1, value=ctx.informe.conclusion)
    celda.font = Font(name="Lato", size=10)
    celda.alignment = Alignment(wrap_text=True, vertical="top")
    hoja.merge_cells(start_row=fila, start_column=1, end_row=fila + 5, end_column=8)
    return fila + 6


def _hoja_controles(libro, ctx):
    hoja = libro.create_sheet("Controles")
    hoja.column_dimensions["A"].width = 8
    hoja.column_dimensions["B"].width = 52
    hoja.column_dimensions["C"].width = 16
    hoja.column_dimensions["D"].width = 90

    fila = _titulo(hoja, 1, "CONTROLES EJECUTADOS")
    fila = _encabezados(hoja, fila, ["COD", "CONTROL", "ESTADO", "DETALLE"])
    for resultado in ctx.resultados:
        fila = _fila(hoja, fila,
                     [resultado.codigo, resultado.nombre,
                      resultado.estado.value, resultado.detalle],
                     relleno=_RELLENO_ESTADO[resultado.estado])
    _nota_alcance(hoja, fila)


def _hoja_liquidacion(libro, ctx):
    hoja = libro.create_sheet("Recalculo")
    for letra, ancho in zip("ABCDEFG", (12, 10, 18, 18, 18, 18, 14)):
        hoja.column_dimensions[letra].width = ancho

    fila = _titulo(hoja, 1, "RECALCULO DE AUDITORIA vs DECLARACION", ancho=7)
    fila = _encabezados(hoja, fila, [
        "ACTIVIDAD", "TARIFA", "BASE AUDITORIA", "BASE DECLARADA",
        "IMPTO AUDITORIA", "IMPTO DECLARADO", "DIFERENCIA"])

    declarado = {a.codigo: a for a in ctx.borrador.actividades}
    for codigo in sorted(ctx.reconstruccion.por_actividad):
        renglon = ctx.reconstruccion.por_actividad[codigo]
        actividad = declarado.get(codigo)
        fila = _fila(hoja, fila, [
            codigo, float(renglon.tarifa),
            float(renglon.base_declarable),
            float(actividad.base) if actividad else "n/a",
            float(renglon.impuesto_declarable),
            float(actividad.impuesto) if actividad else "n/a",
            float(renglon.impuesto_declarable - actividad.impuesto)
            if actividad else "n/a"])

    fila += 1
    fila = _fila(hoja, fila, ["Impuesto exacto contable (sin redondeo)",
                              float(ctx.reconstruccion.total_impuesto_contable)])
    fila = _fila(hoja, fila, ["Impuesto declarable (redondeo al mil)",
                              float(ctx.reconstruccion.total_impuesto_declarable)])


def _hoja_cruces(libro, ctx):
    hoja = libro.create_sheet("Cruces")
    for letra, ancho in zip("ABCDE", (16, 20, 20, 20, 16)):
        hoja.column_dimensions[letra].width = ancho

    fila = _titulo(hoja, 1, "C2 - BALANCE DE PRUEBA vs AUXILIAR 2368", ancho=5)
    fila = _encabezados(hoja, fila, ["CUENTA", "BALANCE", "AUXILIAR", "DIFERENCIA"])
    if ctx.saldos:
        por_cuenta = {}
        for linea in ctx.lineas:
            por_cuenta[linea.cuenta] = por_cuenta.get(linea.cuenta, 0) + float(
                linea.retencion)
        for cuenta in sorted(ctx.saldos):
            balance = float(ctx.saldos[cuenta])
            auxiliar = por_cuenta.get(cuenta, 0)
            fila = _fila(hoja, fila, [cuenta, balance, auxiliar, balance - auxiliar])
    else:
        fila = _fila(hoja, fila, ["NO EJECUTADO: no se obtuvo el balance"])

    fila += 1
    fila = _titulo(hoja, fila, "C4 - RECALCULO POR TERCERO (BASE DEL ERP)", ancho=5)
    fila = _encabezados(hoja, fila, [
        "NIT", "TARIFA", "BASE", "RETENCION CONTABLE", "RECALCULADA"])
    for tercero in ctx.reconstruccion.por_tercero.values():
        fila = _fila(hoja, fila, [
            tercero.nit, float(tercero.tarifa), float(tercero.base),
            float(tercero.retencion_contable),
            float(tercero.retencion_recalculada)])


def _hoja_excepciones(libro, ctx):
    hoja = libro.create_sheet("Excepciones")
    for letra, ancho in zip("ABCDE", (16, 10, 100, 14, 16)):
        hoja.column_dimensions[letra].width = ancho

    fila = _titulo(hoja, 1, "EXCEPCIONES", ancho=5)
    fila = _encabezados(hoja, fila, [
        "SEVERIDAD", "CONTROL", "DESCRIPCION", "RENGLON", "IMPACTO $"])
    for excepcion in ctx.informe.excepciones_ordenadas:
        fila = _fila(hoja, fila, [
            excepcion.severidad.value, excepcion.control, excepcion.descripcion,
            excepcion.renglon, float(excepcion.impacto_pesos)])
    if not ctx.informe.excepciones_ordenadas:
        fila = _fila(hoja, fila, ["Sin excepciones"])
    _nota_alcance(hoja, fila)


def _hoja_parametros(libro, ctx):
    hoja = libro.create_sheet("Parametros")
    hoja.column_dimensions["A"].width = 30
    hoja.column_dimensions["B"].width = 76

    fila = _titulo(hoja, 1, "PARAMETROS Y TRAZABILIDAD", ancho=2)
    fila = _fila(hoja, fila, ["Municipio", ctx.municipio.nombre])
    fila = _fila(hoja, fila, ["Estado de las tarifas",
                              ctx.municipio.estado_tarifas])
    fila = _fila(hoja, fila, [
        "Origen de las cifras",
        "tarifa: cuenta contable | retencion: auxiliar | base: reporte del ERP"])
    fila = _fila(hoja, fila, ["Base derivada de la retencion",
                              "Si" if ctx.reconstruccion.base_es_derivada else "No"])

    fila += 1
    fila = _encabezados(hoja, fila, ["MANIFIESTO DE FUENTES", "VALOR"])
    if ctx.manifiesto is not None:
        fila = _fila(hoja, fila, ["Declarado por", ctx.manifiesto.declarado_por])
        fila = _fila(hoja, fila, ["Fecha de declaracion", ctx.manifiesto.fecha])
        for rol, nombre in sorted(ctx.manifiesto.archivos.items()):
            if isinstance(nombre, list):
                valor = ", ".join(nombre) if nombre else "no obtenido (declarado explicitamente)"
            else:
                valor = nombre if nombre else "no obtenido (declarado explicitamente)"
            fila = _fila(hoja, fila, [rol, valor])
    else:
        fila = _fila(hoja, fila, [
            "Sin manifiesto.json", "carpeta resuelta por convencion de nombres literales"])

    fila += 1
    fila = _encabezados(hoja, fila, ["TARIFA POR CUENTA", "VALOR"])
    for cuenta, tarifa in sorted(ctx.municipio.tarifa_por_cuenta.items()):
        fila = _fila(hoja, fila, [cuenta, float(tarifa)])

    fila += 1
    fila = _encabezados(hoja, fila, ["HUELLA SHA256 DE LAS FUENTES", "VALOR"])
    for clave, valor in sorted(ctx.huellas.items()):
        fila = _fila(hoja, fila, [clave, valor])

    fila += 1
    fila = _encabezados(hoja, fila, ["INSUMOS OBTENIDOS", "ESTADO"])
    for clave in ("borrador", "auxiliar", "balance", "erp", "facturas",
                  "pago_anterior"):
        fila = _fila(hoja, fila,
                     [clave, "obtenido" if clave in ctx.insumos_obtenidos
                      else "NO OBTENIDO"])
    _nota_alcance(hoja, fila)


def _hoja_notas(libro, ctx):
    """5.4: la PRIMERA hoja del papel. Semaforo arriba, una linea por nota,
    ordenadas por importancia, sin parrafos.

    Motivo: las conclusiones largas no se leen. Esta hoja es la puerta de
    entrada, no el reemplazo de la conclusion.
    """
    hoja = libro.active
    hoja.title = "Notas"
    hoja.column_dimensions["A"].width = 6
    hoja.column_dimensions["B"].width = 10
    hoja.column_dimensions["C"].width = 78
    hoja.column_dimensions["D"].width = 20

    luz = evaluar(ctx.informe)
    fila = _titulo(hoja, 1, "LO PRIMERO QUE HAY QUE MIRAR")
    fila = _fila(hoja, fila, ["SOLIDEZ DE LA REVISION", luz.color, luz.motivo])
    fila = _fila(hoja, fila, ["", "", "Califica que tan respaldada esta la "
                                      "revision, no la calidad del borrador."])
    fila += 1

    _SIMBOLOS = {Severidad.HALLAZGO: "[!]", Severidad.OBSERVACION: "[-]",
                 Severidad.AVISO: "[i]"}
    fila = _titulo(hoja, fila, "NOTAS")
    for excepcion in ctx.informe.excepciones_ordenadas:
        impacto = ("impacto %s" % excepcion.impacto_pesos
                   if excepcion.impacto_pesos else "impacto $0")
        fila = _fila(hoja, fila, [_SIMBOLOS[excepcion.severidad],
                                  excepcion.control,
                                  excepcion.descripcion, impacto])

    for codigo in ctx.informe.controles_no_ejecutados:
        resultado = next(r for r in ctx.resultados if r.codigo == codigo)
        fila = _fila(hoja, fila, ["[--]", codigo, resultado.detalle,
                                  "NO EJECUTADO"])
    for codigo in ctx.informe.controles_atestados:
        resultado = next(r for r in ctx.resultados if r.codigo == codigo)
        fila = _fila(hoja, fila, ["[A]", codigo, resultado.detalle, "ATESTADO"])

    if not ctx.informe.excepciones_ordenadas and             not ctx.informe.controles_no_ejecutados:
        fila = _fila(hoja, fila, ["", "", "Sin excepciones ni controles "
                                          "pendientes.", ""])
    return fila


def generar_papel(*argumentos) -> Path:
    """M4: acepta (ctx, ruta) y tambien la forma historica (ruta, ctx).

    El orden original era contraintuitivo y ya provoco una llamada invertida.
    En vez de romper a quien ya lo usa, se resuelve por tipo.
    """
    primero, segundo = argumentos
    if hasattr(primero, "informe"):
        ctx, ruta_salida = primero, segundo
    else:
        ruta_salida, ctx = primero, segundo
    return _generar(ruta_salida, ctx)


def _generar(ruta_salida, ctx) -> Path:
    libro = openpyxl.Workbook()
    _hoja_notas(libro, ctx)
    _hoja_caratula(libro, ctx)
    _hoja_controles(libro, ctx)
    _hoja_liquidacion(libro, ctx)
    _hoja_cruces(libro, ctx)
    _hoja_excepciones(libro, ctx)
    _hoja_parametros(libro, ctx)
    ruta_salida = Path(ruta_salida)
    libro.save(ruta_salida)
    return ruta_salida
