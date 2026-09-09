"""Ingesta del formato historico de ReteICA que lleva el cliente.

Es un libro con una hoja por periodo (2025-01 en adelante) donde la compania
registra lo que declaro cada mes. Sirve para dos cosas distintas:

  - C15: confrontar el periodo revisado contra el borrador que nos entregaron.
    Si el cliente nos pasa un borrador que no coincide con su propio registro,
    eso es un hallazgo antes de mirar la contabilidad.
  - C12: leer el mes anterior. El motor puede leer lo DECLARADO; el PAGO no
    esta aqui y sin el comprobante C12 no se puede ejecutar.

TRAMPA DEL FORMATO: los nombres de hoja estan sucios y no son uniformes.
En el libro real conviven '2026-02 ', '2026- 3', '2026 - 4', '2026 - 05' y
'2026-07'. No basta con f"{anio}-{mes:02d}"; hay que normalizar.
"""

import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

import openpyxl

_RE_PERIODO = re.compile(r"^\s*(\d{4})\s*-\s*(\d{1,2})\s*$")
_MARCA_TOTALES = "B8"
_UNO = Decimal("1")


@dataclass(frozen=True)
class PeriodoHistorico:
    periodo: str
    base_declarada: Decimal
    retenciones_declaradas: Decimal
    hoja: str


def normalizar_periodo(nombre_hoja: str):
    """'2026 - 6' -> '2026-06'. Devuelve None si la hoja no es un periodo."""
    encontrado = _RE_PERIODO.match(str(nombre_hoja).replace("\xa0", " "))
    if not encontrado:
        return None
    anio, mes = encontrado.groups()
    return "%s-%02d" % (anio, int(mes))


def _a_decimal(valor):
    if valor is None:
        return None
    if isinstance(valor, Decimal):
        return valor
    if isinstance(valor, (int, float)):
        return Decimal(str(valor)).quantize(_UNO)
    texto = str(valor).replace(".", "").replace(",", "").strip()
    return Decimal(texto).quantize(_UNO) if texto.lstrip("-").isdigit() else None


def _totales_de(hoja):
    """La fila de totales es la que sigue al marcador 'B8'."""
    filas = list(hoja.iter_rows(values_only=True))
    for indice, fila in enumerate(filas):
        valores = [c for c in fila if c is not None]
        if valores and str(valores[0]).strip().upper().startswith(_MARCA_TOTALES):
            if indice + 1 >= len(filas):
                break
            siguientes = [_a_decimal(c) for c in filas[indice + 1] if c is not None]
            numeros = [v for v in siguientes if v is not None]
            if len(numeros) >= 2:
                return numeros[0], numeros[1]
    return None, None


def leer_formato_historico(ruta: Path) -> dict:
    """periodo normalizado -> PeriodoHistorico."""
    libro = openpyxl.load_workbook(ruta, data_only=True)
    historico = {}
    for nombre in libro.sheetnames:
        periodo = normalizar_periodo(nombre)
        if periodo is None:
            continue
        base, retenciones = _totales_de(libro[nombre])
        if base is None or retenciones is None:
            continue
        historico[periodo] = PeriodoHistorico(
            periodo=periodo, base_declarada=base,
            retenciones_declaradas=retenciones, hoja=nombre)
    return historico


def periodo_anterior(periodo: str) -> str:
    anio, mes = (int(p) for p in periodo.split("-"))
    return "%d-%02d" % (anio - 1, 12) if mes == 1 else "%d-%02d" % (anio, mes - 1)
